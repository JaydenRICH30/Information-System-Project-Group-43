from pathlib import Path
import re

from django.conf import settings
from django.core.files import File
from django.core.management.base import BaseCommand
from django.utils.html import escape

from openpyxl import load_workbook

from login.models import User
from shop.models import Category, Product, ProductImage, Tag


def _normalize_name(value):
    return re.sub(r'[^a-z0-9]+', '', (value or '').lower())


def _build_description(description_en, description_zh):
    blocks = []
    if description_en:
        blocks.append(f'<p>{escape(str(description_en).strip())}</p>')
    if description_zh and str(description_zh).strip() and str(description_zh).strip() != str(description_en).strip():
        blocks.append(f'<p>{escape(str(description_zh).strip())}</p>')
    return ''.join(blocks) or '<p></p>'


class Command(BaseCommand):
    help = 'Import products from the workspace 商品/商品.xlsx file and attach matching images from 商品/商品图片.'

    def add_arguments(self, parser):
        parser.add_argument('--owner-email', dest='owner_email', help='Assign imported products to this merchant email.')

    def handle(self, *args, **options):
        workspace_root = settings.BASE_DIR.parent
        catalog_root = workspace_root / '商品'
        workbook_path = catalog_root / '商品.xlsx'
        image_dir = catalog_root / '商品图片'

        if not workbook_path.exists():
            self.stderr.write(self.style.ERROR(f'Workbook not found: {workbook_path}'))
            return

        if not image_dir.exists():
            self.stderr.write(self.style.ERROR(f'Image directory not found: {image_dir}'))
            return

        owner_email = options.get('owner_email')
        owner = None
        if owner_email:
            owner = User.objects.filter(email=owner_email, role='merchant').first()
            if owner is None:
                self.stderr.write(self.style.ERROR(f'Merchant not found: {owner_email}'))
                return
        if owner is None:
            owner = User.objects.filter(role='merchant').order_by('id').first()
        if owner is None:
            owner = User.objects.create_user(
                username='catalog_import',
                email='catalog.import@example.com',
                password='DemoPass123',
                role='merchant',
                address='Imported from workspace catalog',
            )

        workbook = load_workbook(workbook_path, data_only=True)
        sheet = workbook.active
        rows = list(sheet.iter_rows(values_only=True))
        headers = [str(value).strip() if value is not None else '' for value in rows[0]]

        image_paths = sorted(path for path in image_dir.iterdir() if path.is_file())
        imported_count = 0
        updated_count = 0

        for row in rows[1:]:
            data = dict(zip(headers, row))
            product_name = str(data.get('Product Name') or data.get('商品名称') or '').strip()
            if not product_name:
                continue

            category_name = str(data.get('Category') or data.get('分类') or 'Imported').strip()
            description = _build_description(data.get('Description'), data.get('商品描述'))
            price = data.get('Price (USD)') or 0
            stock = int(data.get('Stock') or 0)
            tags_text = str(data.get('Tags') or '').strip()

            category, _ = Category.objects.get_or_create(name=category_name)
            product, created = Product.objects.get_or_create(
                owner=owner,
                name=product_name,
                defaults={
                    'description': description,
                    'price': price,
                    'stock': stock,
                    'status': 'active',
                    'category': category,
                },
            )
            product.description = description
            product.price = price
            product.stock = stock
            product.status = 'active'
            product.category = category
            product.save()
            product.tags.clear()
            for tag_name in [item.strip() for item in tags_text.split(',') if item.strip()]:
                tag, _ = Tag.objects.get_or_create(name=tag_name)
                product.tags.add(tag)

            normalized_name = _normalize_name(product_name)
            matched_images = [path for path in image_paths if _normalize_name(path.stem).startswith(normalized_name)]
            if matched_images:
                if product.image:
                    product.image.delete(save=False)
                for gallery_image in product.images.all():
                    gallery_image.image.delete(save=False)
                    gallery_image.delete()

                with matched_images[0].open('rb') as handle:
                    product.image.save(matched_images[0].name, File(handle), save=False)
                product.save(update_fields=['image'])

                for gallery_path in matched_images[1:]:
                    with gallery_path.open('rb') as handle:
                        gallery = ProductImage(product=product)
                        gallery.image.save(gallery_path.name, File(handle), save=False)
                        gallery.product = product
                        gallery.save()

            if created:
                imported_count += 1
            else:
                updated_count += 1

        self.stdout.write(self.style.SUCCESS(f'Imported {imported_count} products and updated {updated_count} products for {owner.email}.'))